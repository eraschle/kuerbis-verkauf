"""Import (JSON, eigenes Excel, alte Excel) und Export (JSON, Excel)."""

from __future__ import annotations

import json
import re
import warnings
import zipfile
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .model import Daten, Saison, anzahl_wochen, betrag_normalisieren, daten_aus_dict, statistik, wochentotale
from .storage import speichern

DATEN_KOPF = ["Jahr", "Startdatum", "Datum", "Betrag"]


def importieren(pfad: Path) -> Daten:
    endung = pfad.suffix.lower()
    if endung == ".json":
        try:
            obj = json.loads(pfad.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as e:
            raise ValueError(f"Die JSON-Datei kann nicht gelesen werden: {e}") from None
        return daten_aus_dict(obj)
    if endung in (".xlsx", ".xlsm"):
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                wb = openpyxl.load_workbook(pfad, data_only=True, read_only=True)
        except (OSError, zipfile.BadZipFile, KeyError) as e:
            raise ValueError(f"Die Excel-Datei kann nicht gelesen werden: {e}") from None
        try:
            if "Daten" in wb.sheetnames:
                return _eigenes_excel(wb["Daten"])
            alte = [n for n in wb.sheetnames if re.fullmatch(r"\d\d", n)]
            if alte:
                return _alte_excel(wb, alte)
        finally:
            wb.close()
    raise ValueError("Unbekanntes Dateiformat. Unterstützt: Kürbisverkauf-JSON, Kürbisverkauf-Excel und die alte Excel-Datei.")


def _als_datum(wert) -> date | None:
    if isinstance(wert, datetime):
        return wert.date()
    if isinstance(wert, date):
        return wert
    if isinstance(wert, str):
        try:
            return date.fromisoformat(wert.strip()[:10])
        except ValueError:
            return None
    return None


def _alte_excel(wb, blaetter: list[str]) -> Daten:
    saisons = {}
    for name in blaetter:
        jahr = 1900 + int(name) if int(name) >= 90 else 2000 + int(name)
        zeilen = list(wb[name].iter_rows(min_row=5, max_col=4, values_only=True))
        if not zeilen:
            continue
        start = _als_datum(zeilen[0][2])
        if start is None:
            continue
        eintraege = {}
        for z in zeilen:
            tag = _als_datum(z[2])
            try:
                betrag = betrag_normalisieren(z[3])
            except ValueError:
                raise ValueError(f"Blatt {name}: ungültiger Betrag {z[3]!r} am {tag}") from None
            if tag and betrag is not None:
                eintraege[tag] = betrag
        if eintraege:
            saisons[jahr] = Saison(jahr, start, eintraege)
    return Daten(saisons)


def _eigenes_excel(ws) -> Daten:
    zeilen = ws.iter_rows(values_only=True)
    kopf = [str(x).strip() if x is not None else "" for x in next(zeilen, [])][:4]
    if kopf != DATEN_KOPF:
        raise ValueError("Das Blatt 'Daten' hat nicht die erwarteten Spalten Jahr, Startdatum, Datum, Betrag.")
    saisons: dict[int, Saison] = {}
    for nr, z in enumerate(zeilen, start=2):
        if not z or all(v is None for v in z[:4]):
            continue
        try:
            jahr = int(z[0])
            start = _als_datum(z[1])
            if start is None:
                raise ValueError("Startdatum fehlt")
            s = saisons.setdefault(jahr, Saison(jahr, start, {}))
            tag = _als_datum(z[2])
            betrag = betrag_normalisieren(z[3])
            if tag and betrag is not None:
                s.eintraege[tag] = betrag
        except (TypeError, ValueError) as e:
            raise ValueError(f"Blatt 'Daten', Zeile {nr}: {e}") from None
    return Daten(saisons)


def json_exportieren(pfad: Path, d: Daten) -> None:
    speichern(pfad, d)
    bak = pfad.with_name(pfad.name + ".bak")
    bak.unlink(missing_ok=True)


def excel_exportieren(pfad: Path, d: Daten) -> None:
    wb = openpyxl.Workbook()
    fett = Font(bold=True)

    ws = wb.active
    ws.title = "Daten"
    ws.append(DATEN_KOPF)
    for jahr, s in sorted(d.saisons.items()):
        if not s.eintraege:
            ws.append([jahr, s.start, None, None])  # leere Saison erhalten
        for tag in sorted(s.eintraege):
            ws.append([jahr, s.start, tag, s.eintraege[tag]])
    for zelle in ws[1]:
        zelle.font = fett
    for zeile in ws.iter_rows(min_row=2):
        zeile[1].number_format = zeile[2].number_format = "DD.MM.YYYY"
        zeile[3].number_format = "#,##0.00"
    for i, breite in enumerate([8, 13, 13, 11], start=1):
        ws.column_dimensions[get_column_letter(i)].width = breite
    ws.freeze_panes = "A2"

    ue = wb.create_sheet("Übersicht")
    wochen = max((anzahl_wochen(s) for s in d.saisons.values()), default=16)
    ue.append(["Jahr", "Start"] + [f"Wo {i}" for i in range(1, wochen + 1)] + ["Min", "Max", "Mittel", "Tage", "Total"])
    for jahr, s in sorted(d.saisons.items()):
        w = wochentotale(s)
        w += [None] * (wochen - len(w))
        st = statistik(s)
        ue.append([jahr, s.start] + [x or None for x in w] + [st["min"], st["max"], st["mittel"], st["tage"], st["total"]])
    for zelle in ue[1]:
        zelle.font = fett
    for zeile in ue.iter_rows(min_row=2):
        zeile[1].number_format = "DD.MM.YYYY"
    ue.column_dimensions["B"].width = 12
    ue.freeze_panes = "C2"

    pfad.parent.mkdir(parents=True, exist_ok=True)
    wb.save(pfad)
