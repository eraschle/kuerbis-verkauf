from datetime import date
from pathlib import Path

import pytest

from kuerbis.excel_io import excel_exportieren, importieren, json_exportieren
from kuerbis.model import Daten, Saison, statistik

ALT = Path(__file__).resolve().parents[1] / "Kürbissverkauf 1998-2025.xlsx"


@pytest.fixture(scope="module")
def alt():
    if not ALT.exists():
        pytest.skip("alte Excel-Datei fehlt")
    return importieren(ALT)


def ueb_statistik():
    """Sollwerte aus dem Blatt 'Üb' der alten Excel: {jahr: (min, max, mittel, tage, total)}."""
    import warnings

    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # unlesbare Kopf-/Fusszeile der alten Datei
        wb = openpyxl.load_workbook(ALT, data_only=True, read_only=True)
    zeilen = list(wb["Üb"].iter_rows(values_only=True))
    wb.close()
    kopf = next(i for i, z in enumerate(zeilen) if list(z[:6]) == ["Jahr", "Min", "Max", "Mittel", "Tage", "Total"])
    werte = {}
    for z in zeilen[kopf + 1 :]:
        if not isinstance(z[0], int):
            break
        if z[4]:  # Jahre ohne Verkaufstage auslassen
            werte[z[0]] = z[1:6]
    return werte


def test_alte_excel_stimmt_mit_uebersicht_ueberein(alt):
    soll = ueb_statistik()
    assert soll, "Statistik im Blatt 'Üb' nicht gefunden"
    assert sorted(soll) == sorted(alt.saisons)
    for jahr, (mn, mx, mittel, tage, total) in soll.items():
        st = statistik(alt.saisons[jahr])
        assert st["tage"] == tage, jahr  # Tage mit 0 CHF zählen mit
        assert st["total"] == pytest.approx(total), jahr
        assert st["max"] == pytest.approx(mx), jahr
        assert st["mittel"] == pytest.approx(mittel, abs=0.005), jahr
        assert st["min"] == pytest.approx(mn or 0), jahr  # 'Üb' zeigt 0 als leer


def test_alte_excel_startdatum(alt):
    assert alt.saisons[2025].start == date(2025, 8, 25)
    assert alt.saisons[2014].start == date(2014, 9, 1)


def test_alte_excel_leere_jahre_werden_ausgelassen(alt):
    assert 1998 not in alt.saisons
    assert 1999 not in alt.saisons
    assert min(alt.saisons) == 2000
    assert max(alt.saisons) == 2025


def beispiel():
    return Daten(
        {
            2024: Saison(2024, date(2024, 8, 26), {date(2024, 9, 2): 12.5, date(2024, 9, 3): 40}),
            2025: Saison(2025, date(2025, 8, 25), {}),
        }
    )


def test_excel_rundlauf(tmp_path):
    pfad = tmp_path / "export.xlsx"
    excel_exportieren(pfad, beispiel())
    assert importieren(pfad) == beispiel()


def test_json_rundlauf(tmp_path):
    pfad = tmp_path / "export.json"
    json_exportieren(pfad, beispiel())
    assert importieren(pfad) == beispiel()


def test_alte_excel_rundlauf(alt, tmp_path):
    pfad = tmp_path / "export.xlsx"
    excel_exportieren(pfad, alt)
    assert importieren(pfad) == alt


def test_unbekannte_datei(tmp_path):
    pfad = tmp_path / "x.txt"
    pfad.write_text("hallo", encoding="utf-8")
    with pytest.raises(ValueError):
        importieren(pfad)


def test_fremdes_excel(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Irgendwas"
    pfad = tmp_path / "fremd.xlsx"
    wb.save(pfad)
    with pytest.raises(ValueError):
        importieren(pfad)
